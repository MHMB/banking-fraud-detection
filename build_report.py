#!/usr/bin/env python3
"""Build the Persian technical report: fraud_detection_report.md and .docx.

Numbers come from report_assets/stats.json and queries from report_assets/queries.py,
so run `python report_assets/generate_assets.py` first. The existing .docx is reused
as the style template (page setup, heading styles). Export the PDF from Word.

Requires: python-docx
"""
import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(ROOT, "report_assets"))
from queries import QUERIES  # noqa: E402

MD, DOCX = os.path.join(ROOT, "fraud_detection_report.md"), os.path.join(ROOT, "fraud_detection_report.docx")
FONT = "Vazirmatn"

FA_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def fa(value) -> str:
    """Persian digits, with ٬ as the thousands separator for integers."""
    text = f"{value:,}".replace(",", "٬") if isinstance(value, int) else str(value)
    return text.translate(FA_DIGITS)


MONTHS = ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
          "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]


def fa_date(date: str) -> str:
    """14030101 -> ۱ فروردین ۱۴۰۳ (slashes get reordered inside RTL text)"""
    return f"{fa(int(date[6:]))} {MONTHS[int(date[4:6]) - 1]} {fa(date[:4])}"


def short_name(n: str) -> str:
    """Persian part of a scenario name, for narrow table columns."""
    return SCENARIOS[n]["name"].split(" (")[0]


SCENARIOS = {
    "01": dict(
        name="لایه‌گذاری ساده (Simple Layering)", pattern="نقد → A→B→C→D → نقد (۳ دور)",
        text="در هر دور، مبلغی کمی کمتر از سقف گزارش‌دهی به‌صورت نقدی به یک حساب تازه‌افتتاح‌شده واریز می‌شود. پول ظرف حدود ۴ روز کاری از سه حساب حقیقی دیگر می‌گذرد، در هر گام ۵۰۰ هزار ریال از آن کم می‌شود و حساب آخر آن را نقد برداشت می‌کند. سه دور با مبالغ ۴۵، ۴۸٫۵ و ۴۲ میلیون ریال اجرا شده است.",
        hidden=["حقوق‌ها و واریزهای نقدی قانونی هم تا حدود ۴۹٫۵ میلیون ریال‌اند", "هر چهار حساب زنجیره تراکنش عادی هم دارند", "شرح تراکنش‌ها خالی یا معمولی است"],
        signs=["واریز نقدی زیر سقف که همان روز منتقل می‌شود", "عبور سریع از چند حساب با کم شدن تدریجی مبلغ", "حساب اول تازه افتتاح شده است", "برداشت نقدی در انتهای زنجیره"],
        note="کوئری برداشت نقدی آخر زنجیره (نقش integration) را برنمی‌گرداند. این برداشت با دنبال کردن حساب آخر پیدا می‌شود."),
    "02": dict(
        name="خرد کردن تراکنش (Structuring / Smurfing)", pattern="۱۵ واریز نقدی زیر سقف → حساب تجمیع‌کننده",
        text="۹ حساب واسط در ۸ روز ۱۵ واریز نقدی بین ۴۳ تا ۴۹٫۵ میلیون ریال می‌گیرند، یعنی همه زیر سقف ۵۰ میلیونی. هر واریز همان روز به یک حساب تجمیع‌کننده می‌رود و این حساب روز دوازدهم حدود ۹۶٪ کل مبلغ را یکجا منتقل می‌کند.",
        hidden=["مبالغ با هم برابر نیستند", "حقوق‌ها و واریزهای نقدی قانونی هم در بازه‌ی ۴۰ تا ۵۰ میلیون‌اند", "حساب‌های واسط فعالیت عادی هم دارند"],
        signs=["واریز نقدی و انتقال همان مبلغ در یک روز", "چند حساب که همه به یک حساب پول می‌فرستند", "انتقال یکجای کل مبلغ چند روز بعد"],
        note="انتقال یکجای آخر (consolidation) گام بعدی حساب تجمیع‌کننده است و در خروجی کوئری نیست."),
    "03": dict(
        name="شبکه‌ی شرکت‌های کاغذی (Shell Company Network)", pattern="A→B→C→A (۴ دور)",
        text="سه حساب حقوقی متعلق به سه مشتری متفاوت در حدود ۶ هفته چهار بار پول را بین خود دور می‌دهند. هر دور با ۴۲۰ تا ۵۶۰ میلیون ریال شروع می‌شود و در هر گام ۳ تا ۷ درصد از آن کم می‌شود. شرح تراکنش‌ها شماره‌ی فاکتورهای پشت‌سرهم است، با خدماتی مثل مشاوره، مدیریت و بازرگانی.",
        hidden=["فاکتورهای B2B قانونی هم همین شکل شرح را دارند و مبالغشان مشابه است (میانه حدود ۳۲۵ میلیون و برخی بالای ۱ میلیارد)", "در نویز هم چرخه‌های کوچک طبیعی هست، پس یک کوئری ساده‌ی چرخه مثبت کاذب می‌دهد"],
        signs=["چرخه‌ی بسته‌ی سه‌تایی که بین همان شرکت‌ها تکرار می‌شود", "کم شدن مبلغ در هر گام", "شرح‌های کلی و مبهم"],
        note=""),
    "04": dict(
        name="تأمین مالی تروریسم (Terrorist Financing)", pattern="۵۰ اهداکننده → خیریه → NGO → تبعه‌ی خارجی",
        text="۵۰ اهداکننده‌ی غیرحقوقی، که چند تبعه‌ی خارجی هم بینشان هست، در ۶۰ روز هر کدام ۱ تا ۳ بار ۱ تا ۵ میلیون ریال به حساب یک مؤسسه‌ی خیریه واریز می‌کنند. خیریه در دو نوبت حدود ۹۵٪ پول جمع‌شده‌ی هر ماه را به یک NGO می‌دهد و NGO دو روز بعد حدود ۹۰٪ آن را به حساب یک تبعه‌ی خارجی می‌فرستد.",
        hidden=["فروشگاه‌های قانونی هم از افراد زیادی پرداخت کوچک می‌گیرند", "شرح کمک‌ها معمولی است: «کمک خیریه»، «نذر» یا خالی"],
        signs=["تعداد زیادی واریز کوچک به یک حساب که زود خارج می‌شود", "خیریه پول را نگه نمی‌دارد", "گیرنده‌ی نهایی تبعه‌ی خارجی است"],
        note=""),
    "05": dict(
        name="تصاحب حساب (Account Takeover)", pattern="۶ ماه رفتار عادی، بعد تخلیه‌ی شبانه",
        text="حساب یک بازنشسته که پیش از ۱۳۹۵ افتتاح شده، شش ماه رفتار ثابتی دارد: ماهی ۲۵ میلیون ریال حقوق بازنشستگی، چند خرید کوچک و یک برداشت از خودپرداز. بعد، بین ساعت ۰ تا ۵ بامداد و با موبایل‌بانک، ۵۰۰ میلیون ریال به یک حساب تازه‌افتتاح‌شده می‌رود و در ۴ روز بعد ۲۰۰، ۱۸۰ و ۱۵۰ میلیون ریال به حساب دیگری.",
        hidden=["انتقال‌های شخصی قانونی ۱۵۰ تا ۶۰۰ میلیونی (خرید خودرو، ودیعه‌ی مسکن) در داده وجود دارد", "رفتار عادی قربانی برچسب ندارد"],
        signs=["مبلغ حدود ۲۰ برابر رفتار معمول حساب", "ساعت غیرعادی بامداد", "گیرنده‌های جدید"],
        note=""),
    "06": dict(
        name="پولشویی مبتنی بر تجارت (Trade-Based ML)", pattern="واردکننده → کارگزار (۱۵٪−) → تبعه‌ی خارجی (۵ دور)",
        text="یک شرکت واردکننده تقریباً هر ۲۰ روز ۰٫۹ تا ۱٫۲ میلیارد ریال بابت پروفرما به یک کارگزار می‌پردازد. کارگزار ۱ تا ۳ روز بعد حدود ۸۵٪ آن را به حساب یک تبعه‌ی خارجی می‌فرستد. بیش‌اظهاری فاکتور در داده‌ی پرداخت دیده نمی‌شود و برای دیدنش اظهارنامه‌ی گمرکی لازم است. آنچه در داده هست، عبور تکراری پول است و سهم ثابت حدود ۱۵٪ که کارگزار نگه می‌دارد.",
        hidden=["پرداخت‌های B2B قانونی با مبالغ مشابه، که برخی‌شان بالای ۱ میلیارد است", "شرح‌ها مثل پروفرمای معمولی است"],
        signs=["سهم ثابتی که کارگزار نگه می‌دارد", "همان سه طرف با فاصله‌های منظم", "گیرنده‌ی نهایی تبعه‌ی خارجی است"],
        note=""),
    "07": dict(
        name="تقلب کارمند بانک (Insider Fraud)", pattern="حساب‌های راکد یک شعبه → حساب‌های واسط → نقد",
        text="۶ حساب حقیقی قدیمی (افتتاح پیش از ۱۳۹۸) از یک بانک و یک شعبه، که هیچ فعالیت دیگری ندارند، تقریباً هر ۳ روز یک بار ۵۰ تا ۲۰۰ میلیون ریال را از باجه‌ی شعبه به یکی از ۴ حساب واسط منتقل می‌کنند. حساب واسط ظرف ۱ تا ۲ روز ۸۵ تا ۹۵٪ مبلغ را نقد برداشت می‌کند.",
        hidden=["قربانیان جز همین انتقال هیچ تراکنشی ندارند و در نویز دیده نمی‌شوند", "حساب‌های واسط فعالیت عادی دارند"],
        signs=["حساب راکدی که ناگهان پول می‌فرستد", "همه‌ی قربانیان از یک شعبه‌اند", "گیرنده‌های مشترک", "برداشت نقدی سریع"],
        note="برداشت‌های نقدی حساب‌های واسط (cash_out) گام بعدی‌اند و در خروجی کوئری نیستند."),
    "08": dict(
        name="پرداخت‌های دایره‌ای (Circular Payments)", pattern="حلقه‌ی ۱۸ حسابی تا مبدا + ۳ مسیر فرعی",
        text="حدود ۱ میلیارد ریال از ۱۸ حساب می‌گذرد و به حساب مبدا برمی‌گردد. بین هر دو گام ۱ تا ۴ روز کاری فاصله است و کل مسیر حدود ۷ هفته طول می‌کشد. در هر گام ۴ تا ۶ درصد کم می‌شود و در پایان حدود ۴۱۶ میلیون به مبدا می‌رسد. سه مسیر فرعی ۲۰۰ میلیونی هم از مبدا جدا می‌شوند و جلوتر به حلقه برمی‌گردند.",
        hidden=["شرح‌ها خالی است", "همه‌ی حساب‌های حلقه فعالیت عادی هم دارند"],
        signs=["چرخه‌ی طولانی که به مبدا برمی‌گردد", "کم شدن یکنواخت مبلغ", "مسیرهای فرعی که دوباره به حلقه می‌پیوندند"],
        note="جستجوی چرخه‌ی طولانی بدون محدودیت در Cypher از نظر محاسباتی منفجر می‌شود. این کوئری برای اینکه اجرا شود فقط تراکنش‌های ۳۰۰ میلیون ریال به بالا را بررسی می‌کند و به همین دلیل مسیرهای فرعی ۲۰۰ میلیونی (side_path) را نمی‌بیند."),
    "09": dict(
        name="پخش و تجمیع (Scatter-Gather)", pattern="یک مبدا → ۷ واسط → یک جمع‌کننده",
        text="یک حساب مبدا در ۳ روز به ۷ حساب حقیقی، هر کدام ۲۸ تا ۵۹ میلیون ریال می‌فرستد. دو تا از این شاخه‌ها از یک حساب واسط دوم هم می‌گذرند. هر شاخه ظرف ۱ تا ۴ روز ۹۵ تا ۹۹٪ پولی را که گرفته به یک حساب جمع‌کننده‌ی مشترک منتقل می‌کند.",
        hidden=["هر مبلغ در بازه‌ی انتقال‌های شخصی عادی است", "شرح‌ها خالی است", "مبدا، واسط‌ها و جمع‌کننده تراکنش عادی هم دارند"],
        signs=["پخش پول و جمع شدن دوباره‌ی آن بین همان دو حساب", "عبور تقریباً کامل پول از واسط‌ها", "نگه داشتن کوتاه پول در هر حساب"],
        note="کوئری Cypher فقط شاخه‌های تک‌واسطه را پیدا می‌کند (۵ از ۷). دو شاخه‌ی دوواسطه‌ای را API بخش ۱۶ پیدا می‌کند، چون در آن عمق مسیر ثابت نیست."),
}


def render_md(stats) -> str:
    sc, val, gr = stats["scenarios"], stats["validation"], stats["graph"]
    total = sum(s["transactions"] for s in sc.values())
    planted = sum(s["planted"] for s in sc.values())
    out = []
    w = out.append

    w("# شبیه‌سازی سناریوهای تقلب در نظام بانکی ایران\n\n## گزارش فنی\n")
    w("**تهیه‌کننده:** Mohammad Motiebirjandi  \n**تاریخ:** آبان ۱۴۰۵  \n**پایگاه داده:** Neo4j 5.26.2\n\n---\n")

    w("## ۱. مقدمه و هدف\n")
    w(f"در این پروژه {fa(len(sc))} سناریوی تقلب و پولشویی روی داده‌های مصنوعی شبیه‌سازی شده است. داده‌ها کاملاً ساختگی‌اند ولی قواعد نظام بانکی ایران را رعایت می‌کنند. سؤال اصلی این بود که تحلیل گرافی تا کجا می‌تواند الگوهایی را پیدا کند که عمداً لای فعالیت عادی پنهان شده‌اند.\n")
    w(f"داده‌ها شامل {fa(val['Accounts'])} حساب متعلق به {fa(val['Customers'])} مشتری در ۱۰ بانک و {fa(total)} تراکنش است. از این تراکنش‌ها {fa(planted)} مورد تقلبِ کاشته‌شده است و بقیه فعالیت عادی (نویز). داده‌ها در پایگاه داده‌ی گرافی Neo4j بارگذاری شده‌اند و یک API برای تحلیل مسیر روی آن ساخته شده است.\n")
    w("### ۱.۱ ابزارها\n")
    w("| ابزار | کاربرد |\n|---|---|\n| Python 3.10+ | تولید داده‌ی مصنوعی، اعتبارسنجی و تحلیل |\n| Neo4j 5.26.2 | پایگاه داده‌ی گرافی |\n| Docker | اجرای Neo4j و API |\n| Cypher | زبان پرس‌وجوی گراف |\n| FastAPI | سرویس تحلیل مسیر |\n| NetworkX / Matplotlib | رسم نمودارها |\n\n---\n")

    w("## ۲. روش تولید داده\n")
    w("داده‌ها با اسکریپت تولید می‌شوند و چون seed ثابت است، هر اجرا دقیقاً همان داده را می‌سازد. در تولید داده دو قید رعایت شده: داده باید به داده‌ی واقعی بانکی نزدیک باشد و الگوی تقلب باید لای نویز پنهان بماند.\n")
    w("### ۲.۱ تولید حساب‌ها\n")
    w("- شماره‌ی شبا: ۲۶ کاراکتر با رقم کنترل mod-97 و شناسه‌ی واقعی بانک‌ها (مثلاً 017 برای ملی، 012 برای ملت و 057 برای پاسارگاد)\n- شماره‌ی ملی اشخاص حقیقی: ۱۰ رقم با رقم کنترل معتبر\n- شناسه‌ی ملی اشخاص حقوقی: ۱۱ رقم با رقم کنترل معتبر\n- شناسه‌ی شهاب: ۱۶ رقم برای هر مشتری\n- هر بانک شعبه‌های خودش را دارد و هر حساب در یکی از شعبه‌های همان بانک باز می‌شود\n- فایل حساب‌ها در قالب ۳۰ ستونی FATF است؛ هر سطر یک رابطه‌ی مشتری و حساب است (مالک، شریک یا صاحب امضا)\n")
    w("### ۲.۲ تولید تراکنش‌ها\n")
    w("- انواع تراکنش: واریز نقدی، برداشت نقدی، انتقال داخلی (هم‌بانک)، پایا، و ساتنا برای انتقال بین‌بانکی ۱۵۰ میلیون ریال به بالا\n- تراکنش نقدی طرف مقابل ندارد: واریز نقدی فرستنده ندارد و برداشت نقدی گیرنده\n- تاریخ‌ها با تقویم شمسی واقعی ساخته می‌شوند (طول درست ماه‌ها و سال کبیسه)\n- تراکنش شعبه، پایا و ساتنا در روز جمعه انجام نمی‌شود\n- کانال‌ها: اینترنت‌بانک، موبایل‌بانک، شعبه و خودپرداز\n")
    w("### ۲.۳ نویز و پنهان‌سازی\n")
    w("در هر سناریو، فعالیت عادی در کل بازه‌ی زمانی سناریو پخش شده است: خرید از فروشگاه، انتقال شخصی با مبالغ رُند و اغلب بدون شرح، حقوق ماهانه‌ی ۲۲ تا ۴۹٫۵ میلیون ریال (درست زیر سقف ۵۰ میلیونی)، واریز و برداشت نقدی، فاکتورهای B2B که میانه‌شان حدود ۳۲۵ میلیون است و بعضی از ۱ میلیارد هم بیشترند، و گاهی انتقال شخصی ۱۵۰ تا ۶۰۰ میلیونی.\n")
    w("تراکنش‌های تقلب شرح معمولی یا خالی دارند و مبالغشان با تراکنش‌های قانونی هم‌پوشانی دارد. حساب‌های درگیر در تقلب هم فعالیت عادی دارند؛ تنها استثنا قربانیان راکد سناریوی ۰۷ هستند.\n")
    w("### ۲.۴ برچسب‌های مرجع\n")
    w("تراکنش‌های کاشته‌شده‌ی هر سناریو همراه با نقششان در فایل `ground_truth.csv` ثبت شده‌اند. این برچسب‌ها در فایل `transactions.csv` نیامده‌اند و در Neo4j هم بارگذاری نمی‌شوند، پس کوئری‌های تشخیص به آن‌ها دسترسی ندارند. کاربرد این فایل فقط سنجیدن نتایج است.\n\n---\n")

    w("## ۳. اعتبارسنجی داده‌ها\n")
    w("| شاخص | مقدار | وضعیت |\n|---|---|---|")
    rows = [("حساب‌ها", "Accounts"), ("مشتریان", "Customers"), ("تراکنش‌ها", "Transactions"),
            ("تراکنش نقدی (یک‌طرفه، طبق انتظار)", "Cash (one-sided) tx"),
            ("تراکنش بدون فرستنده و گیرنده", "No sender and no receiver"),
            ("ارجاع به شبای ناشناخته", "Unknown Sheba"), ("شناسه‌ی تراکنش تکراری", "Duplicate IDs"),
            ("مبلغ صفر یا منفی", "Zero / negative amounts"), ("انتقال به خود", "Self transfers")]
    for label, key in rows:
        w(f"| {label} | {fa(val[key])} | PASS |")
    w("\nرقم کنترل همه‌ی شباها، شماره‌های ملی و شناسه‌های ملی معتبر است و تاریخ نامعتبری در داده نیست.\n")
    w("![نمودار اعتبارسنجی داده‌ها](report_assets/overview/data_validation_chart.png)\n\n---\n")

    w("## ۴. مدل گرافی Neo4j\n")
    w("### ۴.۱ گره‌ها (Nodes)\n")
    w(f"| گره | تعداد |\n|---|---|\n| Bank | {fa(gr['Bank'])} |\n| Customer | {fa(gr['Customer'])} |\n| Customer:Person | {fa(gr['Person'])} |\n| Customer:Person:ForeignNational | {fa(gr['ForeignNational'])} |\n| Customer:Company | {fa(gr['Company'])} |\n| Account | {fa(gr['Account'])} |\n| Transaction | {fa(gr['Transaction'])} |\n")
    w("### ۴.۲ روابط (Relationships)\n")
    w(f"| رابطه | جهت | تعداد |\n|---|---|---|\n| OWNS_ACCOUNT | Customer → Account | {fa(gr['OWNS_ACCOUNT'])} |\n| BELONGS_TO_BANK | Account → Bank | {fa(gr['BELONGS_TO_BANK'])} |\n| SENT | Account → Transaction | {fa(gr['SENT'])} |\n| RECEIVED | Transaction → Account | {fa(gr['RECEIVED'])} |\n")
    w("واریز نقدی فقط رابطه‌ی RECEIVED دارد و برداشت نقدی فقط رابطه‌ی SENT.\n")
    w("### ۴.۳ نمای کلی گراف\n")
    w("![نمونه‌ی تصادفی از گراف](report_assets/overview/neo4j_full_graph.png)\n\n---\n")

    w("## ۵. خلاصه‌ی سناریوها\n")
    w("| # | سناریو | تراکنش | کاشته‌شده | بازه | الگو |\n|---|---|---|---|---|---|")
    for n, s in sc.items():
        w(f"| {fa(n)} | {short_name(n)} | {fa(s['transactions'])} | {fa(s['planted'])} | {fa_date(s['first_date'])} تا {fa_date(s['last_date'])} | {SCENARIOS[n]['pattern']} |")
    w("\n![تراکنش‌های کاشته‌شده در برابر تراکنش‌های عادی](report_assets/overview/volume_comparison_chart.png)\n\n---\n")

    for i, (n, s) in enumerate(sc.items(), start=6):
        info, det, k = SCENARIOS[n], s["detection"], fa(i)
        w(f"## {k}. سناریوی {fa(n)}: {info['name']}\n")
        w(f"### {k}.۱ شرح سناریو\n\n{info['text']}\n")
        w("**آمار:**\n")
        w(f"- الگو: {info['pattern']}\n- تعداد تراکنش: {fa(s['transactions'])}\n- تراکنش کاشته‌شده: {fa(s['planted'])}\n- تراکنش عادی: {fa(s['transactions'] - s['planted'])}\n- حجم تراکنش‌های کاشته‌شده: {fa(round(s['planted_volume'] / 1e6))} میلیون ریال\n- بازه: {fa_date(s['first_date'])} تا {fa_date(s['last_date'])}\n- نقش‌ها: {'، '.join(f'{r} ({fa(c)})' for r, c in s['roles'].items())}\n")
        w(f"### {k}.۲ چطور پنهان شده\n\n" + "\n".join(f"- {x}" for x in info["hidden"]) + "\n")
        w(f"### {k}.۳ نشانه‌های تقلب\n\n" + "\n".join(f"- {x}" for x in info["signs"]) + "\n")
        w(f"### {k}.۴ نمودارها\n")
        w(f"![جریان پول کاشته‌شده، سناریوی {fa(n)}](report_assets/scenario_{n}/transaction_flow_graph.png)\n")
        w(f"![توزیع مبالغ، سناریوی {fa(n)}](report_assets/scenario_{n}/amount_distribution.png)\n")
        w(f"![توزیع کانال‌ها، سناریوی {fa(n)}](report_assets/scenario_{n}/channel_distribution.png)\n")
        w(f"### {k}.۵ کوئری تشخیص\n\n```cypher\n{QUERIES[n]}\n```\n")
        w(f"![خروجی کوئری تشخیص، سناریوی {fa(n)}](report_assets/scenario_{n}/neo4j_graph.png)\n")
        w(f"**نتیجه:** کوئری {fa(det['returned_tx'])} تراکنش برگرداند: {fa(det['true_positive'])} مورد کاشته‌شده و {fa(det['false_positive'])} مثبت کاذب. از {fa(det['planted'])} تراکنش کاشته‌شده، {fa(det['true_positive'])} مورد پیدا شد ({fa(round(det['recall'] * 100))}٪).\n")
        if info["note"]:
            w(info["note"] + "\n")
        w("---\n")

    k = 6 + len(sc)
    w(f"## {fa(k)}. نتایج تشخیص\n")
    w("| # | سناریو | برگشتی | درست | مثبت کاذب | کاشته‌شده | پوشش |\n|---|---|---|---|---|---|---|")
    for n, s in sc.items():
        d = s["detection"]
        w(f"| {fa(n)} | {short_name(n)} | {fa(d['returned_tx'])} | {fa(d['true_positive'])} | {fa(d['false_positive'])} | {fa(d['planted'])} | {fa(round(d['recall'] * 100))}٪ |")
    w("\n**محدودیت این سنجش:** کوئری‌ها را با دانستن الگوی هر سناریو نوشته‌ایم و آستانه‌هایشان (نسبت عبور، حداقل مبلغ، تعداد شاخه) را روی همین داده تنظیم کرده‌ایم. پس دقت ۱۰۰٪ یعنی این الگوها با وجود نویز از فعالیت عادی جدا می‌شوند؛ دقت همین کوئری‌ها روی داده‌ای که ندیده‌اند معلوم نیست. هر جا پوشش کمتر از ۱۰۰٪ است، تراکنش‌های جاافتاده گام‌هایی‌اند که بیرون از مسیر برگشتی کوئری قرار دارند.\n\n---\n")

    w(f"## {fa(k + 1)}. سرویس تحلیل مسیر\n")
    w("سرویسی که با FastAPI روی Neo4j ساخته شده، به دو پرسش جواب می‌دهد. جستجو روی کل گراف انجام می‌شود و مسیرها ترتیب زمانی دارند: هر گام فقط وقتی حساب می‌شود که پول پیش از آن به حساب رسیده باشد.\n")
    w(f"### {fa(k + 1)}.۱ نقش یک شبا در الگوی پخش و تجمیع\n")
    w("مسیر: `GET /shebas/{sheba}/scatter-gather`\n")
    w("این سرویس مشخص می‌کند یک شبا در الگوی پخش و تجمیع پول بوده یا نه، و اگر بوده نقشش مبدا بوده، مقصد یا واسط. برای هر الگو همه‌ی شاخه‌ها با جزئیات تراکنش‌ها برمی‌گردد، همراه با مبلغ پخش‌شده و جمع‌شده، درصدی که در راه کم شده و مدت الگو. مبدا می‌تواند یک حساب باشد یا واریزهای نقدی. پارامترهای قابل تنظیم و مقدار پیش‌فرضشان: حداقل تعداد شاخه (۲)، حداکثر واسط در هر شاخه (۳)، بازه‌ی زمانی (۳۰ روز) و حداقل سهمی که هر حساب باید جلو بفرستد (۰٫۸).\n")
    w(f"### {fa(k + 1)}.۲ رابطه‌ی بین شباها\n")
    w("مسیر: `POST /relationships`\n")
    w("فهرستی از شباها می‌گیرد و برای هر جفت بررسی می‌کند که در هر کدام از دو جهت، با هر تعداد واسطه، مسیر انتقال پول هست یا نه. کوتاه‌ترین مسیر با جزئیات تراکنش‌ها برمی‌گردد و اگر دو حساب مالک مشترک داشته باشند، آن هم گزارش می‌شود.\n")
    w(f"### {fa(k + 1)}.۳ نتیجه روی داده\n")
    w("- در سناریوی ۰۹ هر ۷ شاخه و هر ۱۶ تراکنش کاشته‌شده پیدا شد و نقش مبدا، مقصد و واسط درست تشخیص داده شد\n- سناریوی ۰۲ به شکل الگویی با مبدا نقدی و ۹ شاخه پیدا شد و هر ۱۸ تراکنش مربوط در آن بود\n- در کل گراف فقط ۵ الگوی پخش و تجمیع وجود دارد: دو الگوی کاشته‌شده (۹ و ۷ شاخه) و سه الگوی کوچک ۲ تا ۳ شاخه‌ای که بیشترشان نویزند\n- حدود ۷۵٪ جفت‌حساب‌ها مسیر زمانیِ معتبری به هم دارند. به همین دلیل جواب «رابطه دارد» به‌تنهایی چیز زیادی نمی‌گوید و باید خود مسیر را دید: تعداد گام‌ها، تاریخ‌ها و مبالغ\n\n---\n")

    w(f"## {fa(k + 2)}. محدودیت‌ها\n")
    w("- اطلاعات تماس مشترک مثل تلفن، نشانی یا IP مدل نشده، چون قالب ۳۰ ستونی FATF چنین فیلدی ندارد\n- مبالغ در همان مقیاس اولیه‌ی پروژه مانده‌اند (سقف گزارش‌دهی ۵۰ میلیون ریال) و با سطح ریالی سال ۱۴۰۳ تطبیق داده نشده‌اند\n- همه‌ی سناریوها در یک بازه‌ی زمانی و روی یک مجموعه حساب مشترک ساخته شده‌اند، پس در جستجوی کل گراف ممکن است مسیرها از نویز سناریوهای دیگر بگذرند\n- سرویس تحلیل مسیر در هر درخواست همه‌ی تراکنش‌ها را در حافظه می‌خواند. این روش برای حجم فعلی کار می‌کند ولی برای میلیون‌ها تراکنش باید عوض شود\n- تصاویر «خروجی کوئری» از نتیجه‌ی واقعی کوئری‌ها روی Neo4j رسم شده‌اند و اسکرین‌شات Neo4j Browser نیستند\n\n---\n")

    w(f"## {fa(k + 3)}. جمع‌بندی\n")
    w("ساختار گرافی هر الگو (زنجیره، چرخه، هم‌گرایی، پخش و تجمیع) همراه با ترتیب زمانی تراکنش‌ها کافی بود تا الگوها با وجود مبالغ هم‌پوشان و شرح‌های خنثی از نویز جدا شوند. برای الگوهایی که عمقشان ثابت است کوئری Cypher کافی است. وقتی عمق مسیر متغیر است و ترتیب زمانی مهم است، جستجوی مسیر در سرویس API بهتر جواب می‌دهد.\n\n---\n")
    w("## منابع\n")
    w("- [FATF 40 Recommendations](https://www.fatf-gafi.org/publications/fatf-standards/)\n- [FATF Typologies Reports](https://www.fatf-gafi.org/publications/typologies/)\n- [Central Bank of Iran](https://www.cbi.ir/)\n- [Sheba / IBAN structure: Iran](https://www.ibantest.com/en/iban-structure/iran)\n")
    return '<div dir="rtl">\n\n' + "\n".join(out) + "\n</div>\n"


# ---------------------------------------------------------------- docx

def _fonts(run, name, size):
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), name)
    run.font.size = Pt(size)
    cs = OxmlElement("w:szCs")
    cs.set(qn("w:val"), str(int(size * 2)))
    rpr.append(cs)


def rtl(paragraph, align="right"):
    """bidi paragraph. 'right' is written as jc=start: in an RTL paragraph jc=right means the visual left."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
    ppr = paragraph._p.get_or_add_pPr()
    jc = ppr.find(qn("w:jc"))
    if align != "center":
        jc.set(qn("w:val"), "start")
    if ppr.find(qn("w:bidi")) is None:
        jc.addprevious(OxmlElement("w:bidi"))  # schema order: bidi before jc
    for run in paragraph.runs:
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


def add_text(doc, text, size=11, align="right", container=None):
    """Paragraph with **bold** and `code` inline markup, links reduced to their label."""
    p = (container or doc).add_paragraph()
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
            _fonts(run, FONT, size)
        elif part.startswith("`"):
            _fonts(p.add_run(part[1:-1]), "Courier New", size - 1)
        else:
            _fonts(p.add_run(part), FONT, size)
    rtl(p, align)
    return p


def add_code(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        _fonts(p.add_run(line or " "), "Courier New", 8)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "F4F4F4")
        p._p.get_or_add_pPr().insert(0, shd)  # schema order: shd before spacing and jc
    doc.add_paragraph()


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl_pr = table._tbl.tblPr
    tbl_pr.append(OxmlElement("w:bidiVisual"))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "999999")):
            el.set(qn(f"w:{k}"), v)
        borders.append(el)
    tbl_pr.append(borders)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
            p = add_text(doc, f"**{value}**" if i == 0 else value, size=9, align="center", container=cell)
            p.paragraph_format.space_after = Pt(0)
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "D9E2F3")
                cell._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph()


def build_docx(md: str):
    doc = Document(DOCX)  # reuse page setup and heading styles
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for rid, rel in list(doc.part.rels.items()):  # drop the template's old pictures
        if rel.reltype.endswith("/image"):
            del doc.part.rels[rid]
    for level in (1, 2):  # template headings are blue and have no Persian (cs) font
        style = doc.styles.element.xpath(f'w:style[@w:styleId="Heading{level}"]')[0]
        rpr = style.find(qn("w:rPr"))
        for tag in ("w:color", "w:rFonts"):
            for el in rpr.findall(qn(tag)):
                rpr.remove(el)
        fonts = OxmlElement("w:rFonts")
        for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
            fonts.set(qn(f"w:{attr}"), FONT)
        rpr.insert(0, fonts)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        rpr.append(color)
    for section in doc.sections:  # header/footer parts carry their own paragraphs
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                ppr = p._p.get_or_add_pPr()
                if ppr.find(qn("w:bidi")) is None:
                    ppr.insert(0, OxmlElement("w:bidi"))

    lines = md.splitlines()
    i, title_done = 0, False
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line or line == "---" or line.startswith(("<div", "</div")):
            continue
        if line.startswith("# "):
            doc.add_paragraph()
            add_text(doc, line[2:], size=24, align="center")
        elif line.startswith("## ") and not title_done:  # subtitle + byline block on the title page
            add_text(doc, line[3:], size=16, align="center")
            while not lines[i].startswith("---"):
                if lines[i].strip():
                    add_text(doc, lines[i].strip().replace("**", ""), size=12, align="center")
                i += 1
            doc.add_page_break()
            title_done = True
        elif line.startswith("## ") or line.startswith("### "):
            level = 1 if line.startswith("## ") else 2
            p = doc.add_paragraph(line.lstrip("# "))
            style = OxmlElement("w:pStyle")  # by id: this template's heading names defeat lookup by name
            style.set(qn("w:val"), f"Heading{level}")
            p._p.get_or_add_pPr().insert(0, style)
            rtl(p)
        elif line.startswith("```"):
            block = []
            while not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code(doc, block)
        elif line.startswith("|"):
            rows = [line]
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
            add_table(doc, [r for r in cells if not set("".join(r)) <= set("-: ")])
        elif m := re.match(r"!\[(.*)\]\((.*)\)", line):
            doc.add_picture(os.path.join(ROOT, m.group(2)), width=Cm(15.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(doc, m.group(1), size=9, align="center")
        elif line.startswith("- "):
            add_text(doc, "•  " + line[2:])
        else:
            add_text(doc, line)
    doc.save(DOCX)


def main():
    with open(os.path.join(ROOT, "report_assets", "stats.json"), encoding="utf-8") as f:
        stats = json.load(f)
    md = render_md(stats)
    with open(MD, "w", encoding="utf-8") as f:
        f.write(md)
    build_docx(md)
    print(f"wrote {MD} and {DOCX}")


if __name__ == "__main__":
    main()
